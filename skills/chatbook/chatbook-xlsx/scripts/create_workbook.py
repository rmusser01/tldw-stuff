# SPDX-License-Identifier: Apache-2.0
# Copyright 2026 Robert Benjamin Jake Musser
"""Small fictional cost workbook. Replace sample inputs for a real request."""

import argparse
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.datavalidation import DataValidation


def build(output: Path) -> None:
    """Write a formula-based example without overwriting an existing file."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Costs"
    for row in [
        ["Item", "Quantity", "Unit cost (USD)", "Total (USD)", "Source"],
        ["Pilot materials", 2, 12, "=B2*C2", "Fictional example"],
        ["Review sessions", 3, 7, "=B3*C3", "Fictional example"],
        ["Grand total", None, None, "=SUM(D2:D3)", None],
    ]:
        sheet.append(row)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="123A52")
    for row in sheet.iter_rows(min_row=2, max_row=3, min_col=2, max_col=3):
        for cell in row:
            cell.fill = PatternFill("solid", fgColor="E2F1FA")
    for row in sheet.iter_rows(min_row=2, max_row=4, min_col=3, max_col=4):
        for cell in row:
            cell.number_format = '"$"#,##0.00'
    for column, width in {"A": 24, "B": 14, "C": 20, "D": 20, "E": 26}.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = "A1:E3"
    validation = DataValidation(
        type="decimal", operator="greaterThanOrEqual", formula1=0
    )
    validation.showErrorMessage = True
    validation.error = "Enter a non-negative quantity."
    sheet.add_data_validation(validation)
    validation.add("B2:B3")
    chart = BarChart()
    chart.title = "Example line costs"
    chart.y_axis.title = "USD"
    chart.add_data(
        Reference(sheet, min_col=4, min_row=1, max_row=3), titles_from_data=True
    )
    chart.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=3))
    sheet.add_chart(chart, "A7")
    notes = workbook.create_sheet("Notes")
    notes.append(["Purpose", "Fictional example; not financial evidence."])
    notes.append(["Inputs", "Blue cells are editable; currency is USD."])
    notes.append(
        ["Calculation", "Results require recalculation by Excel or LibreOffice."]
    )
    notes.column_dimensions["A"].width = 18
    notes.column_dimensions["B"].width = 76
    workbook.calculation = CalcProperties(fullCalcOnLoad=True)
    with output.open("xb") as stream:
        workbook.save(stream)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "output", type=Path, help="New .xlsx path in an existing directory"
    )
    args = parser.parse_args()
    build(args.output)
