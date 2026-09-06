# SPDX-License-Identifier: Apache-2.0
# Copyright 2026 Robert Benjamin Jake Musser
"""Small fictional project-update deck. Adapt its content for a real request."""

import argparse
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt


def build(output: Path) -> None:
    """Create three editable example slides without replacing existing files."""
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    slides = [
        (
            "Project update",
            ["Fictional example", "Replace these statements with sourced material."],
            "Audience: project team. No measured results are supplied.",
        ),
        (
            "Pilot before expansion",
            [
                "Define a small, reversible pilot.",
                "Agree on success criteria before starting.",
            ],
            "Recommendation for this fictional example; it is not a reported decision.",
        ),
        (
            "Resolve ownership and timing",
            ["Pilot owner: unresolved", "Decision date: unresolved"],
            "Ask the team to supply the missing owner and decision date.",
        ),
    ]
    for title, lines, notes in slides:
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = title
        title_run = slide.shapes.title.text_frame.paragraphs[0].runs[0]
        title_run.font.size = Pt(32)
        title_run.font.color.rgb = RGBColor.from_string("123A52")
        frame = slide.placeholders[1].text_frame
        frame.clear()
        for index, text in enumerate(lines):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.text = text
            paragraph.font.size = Pt(24)
            paragraph.space_after = Pt(16)
        slide.notes_slide.notes_text_frame.text = notes
    with output.open("xb") as stream:
        presentation.save(stream)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "output", type=Path, help="New .pptx path in an existing directory"
    )
    args = parser.parse_args()
    build(args.output)
