# SPDX-License-Identifier: Apache-2.0
# Copyright 2026 Robert Benjamin Jake Musser
"""Small fictional decision-brief example. Adapt its content for a real request."""

import argparse
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt


def build(output: Path) -> None:
    """Create an editable example without overwriting an existing file."""
    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.8)
    section.left_margin = section.right_margin = Inches(0.8)
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(7)
    document.add_heading("Project decision brief", level=0)
    document.add_paragraph("Fictional example — replace with verified source material.")
    document.add_heading("Recommendation", level=1)
    paragraph = document.add_paragraph()
    paragraph.add_run("Draft decision").bold = True
    paragraph.add_run(": run a small pilot before expanding the service.")
    document.add_heading("Evidence and open questions", level=1)
    document.add_paragraph(
        "No measured results have been supplied.", style="List Bullet"
    )
    document.add_paragraph(
        "Confirm the pilot scope and success criteria.", style="List Bullet"
    )
    document.add_heading("Actions", level=1)
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, ("Action", "Owner", "Due")):
        cell.text = text
        cell.paragraphs[0].runs[0].bold = True
    for cell, text in zip(
        table.add_row().cells, ("Define pilot criteria", "Unresolved", "Unresolved")
    ):
        cell.text = text
    with output.open("xb") as stream:
        document.save(stream)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "output", type=Path, help="New .docx path in an existing directory"
    )
    args = parser.parse_args()
    build(args.output)
